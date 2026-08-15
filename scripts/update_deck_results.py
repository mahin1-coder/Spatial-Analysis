#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs_all_cases" / "presentation" / "tornado_damage_path_analysis_all_29_cases_multipath_DAT_partial.pptx"
OUTPUT = ROOT / "outputs_all_cases" / "presentation" / "tornado_damage_path_analysis_all_29_cases_water_direction_fixed.pptx"
RESULTS = Path("/private/tmp/tornado_outputs_all_cases/cases")
ASSETS = Path("/private/tmp/tornado_slide_assets")
CASES = [
    "TOR5", "TOR7", "TOR10", "TOR11", "TOR12", "TOR13", "TOR15", "TOR16",
    "TOR18", "TOR24", "TOR61", "TOR62", "TOR66", "TOR68", "TOR69", "TOR70",
    "TOR77", "TOR78", "TOR90", "TOR91", "TOR95", "TOR101", "TOR102", "TOR105",
    "TOR111", "TOR112", "TOR114", "TOR115", "TOR123",
]


def replace_largest_picture(slide, image_path: Path) -> None:
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pictures:
        raise RuntimeError(f"No picture found on slide for {image_path}")
    picture = max(pictures, key=lambda shape: shape.width * shape.height)
    left, top, width, height = picture.left, picture.top, picture.width, picture.height
    picture._element.getparent().remove(picture._element)
    slide.shapes.add_picture(str(image_path), left, top, width, height)


def update_status(slide, metrics: dict) -> None:
    predicted = int(metrics.get("predicted_path_count", 0))
    matched = int(metrics.get("matched_path_count", 0))
    reference = int(metrics.get("reference_path_count", 0))
    unverified = metrics.get("unverified_predicted_path_count")
    unverified = predicted if unverified is None else int(unverified)
    if reference:
        text = (
            f"{predicted} imagery candidate(s) | {matched}/{reference} available DAT path(s) recovered | "
            f"{unverified} additional candidate(s) unverified"
        )
    else:
        text = (
            f"{predicted} imagery candidate(s) | no DAT path in the available layer | "
            "all candidates unverified"
        )
    for shape in slide.shapes:
        if hasattr(shape, "text") and "imagery candidate(s)" in shape.text:
            shape.text_frame.paragraphs[0].runs[0].text = text
            return
    raise RuntimeError("Result status textbox not found")


def main() -> int:
    presentation = Presentation(SOURCE)
    for index, case_id in enumerate(CASES):
        model_slide = presentation.slides[3 + index * 9 + 7]
        result_slide = presentation.slides[3 + index * 9 + 8]
        replace_largest_picture(model_slide, ASSETS / case_id / "model_prediction_panel.jpg")
        replace_largest_picture(result_slide, ASSETS / case_id / "model_final_path_map.jpg")
        metrics = json.loads((RESULTS / case_id / "model_case_metrics.json").read_text())
        update_status(result_slide, metrics)
    presentation.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
